"""以确定性规则从原始企业表构建高可信企业参考库。"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook


REQUIRED_HEADERS = ("dwmc", "dwzzjgdm", "dwxz", "dwhy", "gzzwlb")
FINAL_HEADERS = ("dwmc", "dwzzjgdm", "dwxz", "dwhy")
PLACEHOLDER = "30"
ACCEPTANCE_REFERENCE_COUNT = 4356
_TRIMMABLE = " \t\r\n\v\f\u00a0\u200b\u2060\ufeff"


@dataclass(frozen=True)
class SourceRecord:
    row_number: int
    dwmc: str
    dwzzjgdm: str
    dwxz: str
    dwhy: str


@dataclass(frozen=True)
class BuildOutcome:
    trusted_rows: list[dict[str, str]]
    removed_by_code: dict[str, dict[str, Any]]
    invalid_code_records: list[SourceRecord]
    statistics: dict[str, int]


def normalize_text(value: Any) -> str:
    """仅清除字段两端的可见或不可见空白，不改变企业语义。"""
    if value is None:
        return ""
    return str(value).strip(_TRIMMABLE)


def normalize_credit_code(value: Any) -> str:
    """信用代码统一为去空白的大写字符串，不进行任何名称或代码猜测。"""
    text = normalize_text(value)
    return "".join(character for character in text if not character.isspace()).upper()


def is_missing_reference_value(value: Any) -> bool:
    """空白和历史占位符 30 都不是真实参考值。"""
    return normalize_text(value) in {"", PLACEHOLDER}


def _to_record(row: dict[str, Any]) -> SourceRecord:
    return SourceRecord(
        row_number=int(row["row_number"]),
        dwmc=normalize_text(row.get("dwmc")),
        dwzzjgdm=normalize_credit_code(row.get("dwzzjgdm")),
        dwxz=normalize_text(row.get("dwxz")),
        dwhy=normalize_text(row.get("dwhy")),
    )


def _valid_values(records: Iterable[SourceRecord], field: str) -> set[str]:
    return {value for value in (getattr(record, field) for record in records) if not is_missing_reference_value(value)}


def _detail(records: list[SourceRecord], reasons: set[str]) -> dict[str, Any]:
    return {
        "names": sorted(_valid_values(records, "dwmc")),
        "natures": sorted(_valid_values(records, "dwxz")),
        "industries": sorted(_valid_values(records, "dwhy")),
        "row_numbers": sorted(record.row_number for record in records),
        "record_count": len(records),
        "reasons": sorted(reasons),
    }


def build_trusted_reference(raw_rows: Iterable[dict[str, Any]]) -> BuildOutcome:
    """构建一企业一行的高可信库；任何冲突均整组剔除。"""
    records = [_to_record(row) for row in raw_rows]
    valid_code_records = [record for record in records if not is_missing_reference_value(record.dwzzjgdm)]
    invalid_code_records = [record for record in records if is_missing_reference_value(record.dwzzjgdm)]
    by_code: dict[str, list[SourceRecord]] = defaultdict(list)
    for record in valid_code_records:
        by_code[record.dwzzjgdm].append(record)

    reasons_by_code: dict[str, set[str]] = defaultdict(set)
    code_details: dict[str, dict[str, Any]] = {}
    name_conflict_codes = 0
    nature_conflict_codes = 0
    industry_conflict_codes = 0
    missing_core_codes = 0
    for code, group in by_code.items():
        names = _valid_values(group, "dwmc")
        natures = _valid_values(group, "dwxz")
        industries = _valid_values(group, "dwhy")
        if len(names) > 1:
            reasons_by_code[code].add("同一统一社会信用代码对应多个单位名称")
            name_conflict_codes += 1
        if len(natures) > 1:
            reasons_by_code[code].add("同一企业历史单位性质不一致")
            nature_conflict_codes += 1
        if len(industries) > 1:
            reasons_by_code[code].add("同一企业历史单位行业不一致")
            industry_conflict_codes += 1
        missing_any = False
        if not names:
            reasons_by_code[code].add("单位名称无有效历史参考值")
            missing_any = True
        if not natures:
            reasons_by_code[code].add("单位性质无有效历史参考值")
            missing_any = True
        if not industries:
            reasons_by_code[code].add("单位行业无有效历史参考值")
            missing_any = True
        if missing_any:
            missing_core_codes += 1
        code_details[code] = _detail(group, reasons_by_code[code])

    # 反向检查使用所有有有效代码且有有效名称的记录；不猜测哪一个代码正确。
    codes_by_name: dict[str, set[str]] = defaultdict(set)
    for record in valid_code_records:
        if not is_missing_reference_value(record.dwmc):
            codes_by_name[record.dwmc].add(record.dwzzjgdm)
    conflicting_names = {name: codes for name, codes in codes_by_name.items() if len(codes) > 1}
    reverse_conflict_codes: set[str] = set()
    for codes in conflicting_names.values():
        reverse_conflict_codes.update(codes)
    for code in reverse_conflict_codes:
        reasons_by_code[code].add("同一单位名称对应多个统一社会信用代码")

    removed_by_code = {
        code: _detail(by_code[code], reasons_by_code[code])
        for code in sorted(by_code)
        if reasons_by_code[code]
    }
    trusted_rows = []
    for code in sorted(by_code):
        if code in removed_by_code:
            continue
        detail = _detail(by_code[code], set())
        trusted_rows.append(
            {
                "dwmc": detail["names"][0],
                "dwzzjgdm": code,
                "dwxz": detail["natures"][0],
                "dwhy": detail["industries"][0],
            }
        )

    duplicate_counts = Counter((record.dwmc, record.dwzzjgdm, record.dwxz, record.dwhy) for record in records)
    statistics = {
        "source_rows": len(records),
        "valid_credit_code_records": len(valid_code_records),
        "invalid_credit_code_records": len(invalid_code_records),
        "unique_valid_credit_codes": len(by_code),
        "unique_valid_company_names": len(codes_by_name),
        "complete_duplicate_historical_records": sum(count - 1 for count in duplicate_counts.values() if count > 1),
        "credit_code_multiple_names": name_conflict_codes,
        "company_names_multiple_credit_codes": len(conflicting_names),
        "codes_removed_due_to_name_multiple_codes": len(reverse_conflict_codes),
        "nature_conflict_codes": nature_conflict_codes,
        "industry_conflict_codes": industry_conflict_codes,
        "multiple_conflict_codes": sum(1 for detail in removed_by_code.values() if len(detail["reasons"]) > 1),
        "missing_core_field_codes": missing_core_codes,
        "removed_enterprises": len(removed_by_code),
        "trusted_enterprises": len(trusted_rows),
        "trusted_historical_records": sum(len(by_code[row["dwzzjgdm"]]) for row in trusted_rows),
    }
    return BuildOutcome(trusted_rows, removed_by_code, invalid_code_records, statistics)


def validate_final_reference(rows: list[dict[str, str]]) -> None:
    """独立断言最终文件满足一企业一行及四字段完整性。"""
    codes = [row["dwzzjgdm"] for row in rows]
    assert len(codes) == len(set(codes)), "最终参考库存在重复统一社会信用代码"
    names_to_codes: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        for field in FINAL_HEADERS:
            assert not is_missing_reference_value(row[field]), f"最终参考库字段 {field} 存在缺失或占位符"
        names_to_codes[row["dwmc"]].add(row["dwzzjgdm"])
    assert all(len(codes) == 1 for codes in names_to_codes.values()), "最终参考库存在名称对应多个信用代码"


def validate_final_workbook(path: Path) -> None:
    """重新打开已生成文件，独立检查实际落盘内容而非内存对象。"""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        headers = tuple(normalize_text(value) for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        assert headers == FINAL_HEADERS, "最终参考库列结构不符合 dwmc/dwzzjgdm/dwxz/dwhy"
        rows = [dict(zip(FINAL_HEADERS, (normalize_text(value) for value in values))) for values in sheet.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()
    validate_final_reference(rows)


def read_source_rows(input_path: Path) -> list[dict[str, Any]]:
    """只读读取源表，按表头定位字段，绝不改写输入文件。"""
    workbook = load_workbook(input_path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        headers = [normalize_text(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        missing = [header for header in REQUIRED_HEADERS if header not in headers]
        if missing:
            raise ValueError(f"输入 Excel 缺少必要字段：{', '.join(missing)}")
        positions = {header: headers.index(header) for header in REQUIRED_HEADERS}
        return [
            {
                "row_number": row_number,
                **{header: values[index] for header, index in positions.items()},
            }
            for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2)
        ]
    finally:
        workbook.close()


def _write_workbook(path: Path, headers: tuple[str, ...], rows: Iterable[Iterable[Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    workbook.close()


def _report_text(input_path: Path, outcome: BuildOutcome) -> str:
    stats = outcome.statistics
    lines = [
        "企业参考库高可信清洗报告",
        f"输入文件：{input_path}",
        f"原始数据总行数：{stats['source_rows']}",
        f"有效信用代码记录数：{stats['valid_credit_code_records']}",
        f"无有效信用代码记录数：{stats['invalid_credit_code_records']}",
        f"不同有效信用代码数量：{stats['unique_valid_credit_codes']}",
        f"不同有效单位名称数量：{stats['unique_valid_company_names']}",
        f"完全重复历史记录数：{stats['complete_duplicate_historical_records']}",
        f"信用代码对应多个单位名称的企业数量：{stats['credit_code_multiple_names']}",
        f"单位名称对应多个信用代码的名称数量：{stats['company_names_multiple_credit_codes']}",
        f"因同名多代码涉及并剔除的信用代码数量：{stats['codes_removed_due_to_name_multiple_codes']}",
        f"单位性质存在冲突的企业数量：{stats['nature_conflict_codes']}",
        f"单位行业存在冲突的企业数量：{stats['industry_conflict_codes']}",
        f"同时存在多个剔除原因的企业数量：{stats['multiple_conflict_codes']}",
        f"因核心参考字段无有效值而剔除的企业数量：{stats['missing_core_field_codes']}",
        f"最终被剔除企业总数（有效信用代码去重后）：{stats['removed_enterprises']}",
        f"最终高可信企业数量：{stats['trusted_enterprises']}",
        f"高可信企业对应的历史记录数量：{stats['trusted_historical_records']}",
        f"验收参考值：{ACCEPTANCE_REFERENCE_COUNT}",
        f"与验收参考值差异：{stats['trusted_enterprises'] - ACCEPTANCE_REFERENCE_COUNT}",
        "说明：4356 仅用于差异诊断，未参与任何筛选或保留规则。",
    ]
    return "\n".join(lines) + "\n"


def write_outputs(input_path: Path, output_path: Path, outcome: BuildOutcome, force: bool = False) -> dict[str, Path]:
    """输出高可信库、剔除明细、报告；默认拒绝覆盖既有文件。"""
    directory = output_path.parent
    paths = {
        "trusted": output_path,
        "removed": directory / "企业参考库_剔除明细.xlsx",
        "report": directory / "企业参考库_清洗报告.txt",
    }
    if outcome.statistics["trusted_enterprises"] != ACCEPTANCE_REFERENCE_COUNT:
        paths["difference"] = directory / "企业参考库_数量差异诊断.xlsx"
    if input_path.resolve() in {path.resolve() for path in paths.values()}:
        raise ValueError("输出文件不能覆盖输入文件")
    existing = [path for path in paths.values() if path.exists()]
    if existing and not force:
        raise FileExistsError("输出文件已存在；如确认覆盖，请显式使用 --force：" + ", ".join(str(path) for path in existing))
    directory.mkdir(parents=True, exist_ok=True)
    _write_workbook(paths["trusted"], FINAL_HEADERS, ([row[header] for header in FINAL_HEADERS] for row in outcome.trusted_rows))
    detail_headers = ("统一社会信用代码", "涉及单位名称", "涉及单位性质", "涉及单位行业", "历史记录数量", "剔除原因", "原始Excel行号")
    detail_rows = [
        (code, "；".join(detail["names"]), "；".join(detail["natures"]), "；".join(detail["industries"]), detail["record_count"], "；".join(detail["reasons"]), ",".join(map(str, detail["row_numbers"])))
        for code, detail in outcome.removed_by_code.items()
    ]
    detail_rows.extend(
        ("", record.dwmc, record.dwxz, record.dwhy, 1, "无有效统一社会信用代码", record.row_number)
        for record in outcome.invalid_code_records
    )
    _write_workbook(paths["removed"], detail_headers, detail_rows)
    paths["report"].write_text(_report_text(input_path, outcome), encoding="utf-8")
    if "difference" in paths:
        _write_workbook(paths["difference"], detail_headers, detail_rows)
    return paths


def main() -> None:
    parser = argparse.ArgumentParser(description="构建高可信企业参考库（只读源 Excel）。")
    parser.add_argument("--input", required=True, type=Path, help="原始企业参考表 .xlsx")
    parser.add_argument("--output", type=Path, help="高可信参考库输出路径")
    parser.add_argument("--force", action="store_true", help="明确允许覆盖已有输出文件")
    args = parser.parse_args()
    input_path = args.input.resolve()
    if not input_path.is_file() or input_path.suffix.lower() != ".xlsx":
        raise ValueError("--input 必须是存在的 .xlsx 文件")
    output_path = (args.output or input_path.with_name("企业参考库_高可信.xlsx")).resolve()
    raw_rows = read_source_rows(input_path)
    outcome = build_trusted_reference(raw_rows)
    validate_final_reference(outcome.trusted_rows)
    paths = write_outputs(input_path, output_path, outcome, args.force)
    validate_final_workbook(paths["trusted"])
    print(_report_text(input_path, outcome), end="")
    for label, path in paths.items():
        print(f"{label}：{path}")


if __name__ == "__main__":
    main()
