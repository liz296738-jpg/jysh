from __future__ import annotations

from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from shutil import copy2
import warnings

from openpyxl import load_workbook

from src.audit_workbook import RESULT_HEADERS, audit_workbook
from src.config import FIELD_MAP
from src.reference import ReferenceRecord, write_reference_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SAMPLES = PROJECT_ROOT / "samples"


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _reference(conflict: bool = False) -> ReferenceRecord:
    return ReferenceRecord(
        credit_code="91310000TEST00001X",
        company_name="匿名测试有限公司",
        company_type="民营企业",
        company_type_code="300",
        industry="信息传输、软件和信息技术服务业",
        industry_code="I",
        reference_conflict=conflict,
        conflict_fields=("company_name",) if conflict else (),
    )


def _prepare_input(tmp_path: Path, name: str, **changes: str) -> Path:
    input_path = tmp_path / name
    copy2(SAMPLES / "初审样例.xlsx", input_path)
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style.*")
        workbook = load_workbook(input_path)
    sheet = workbook.active
    header_columns = {cell.value: cell.column for cell in sheet[1]}
    values = {
        "xm": "匿名学生",
        "xh": "ANON-001",
        "dwmc": "匿名测试有限公司",
        "dwzzjgdm": "91310000TEST00001X",
        "dwxz": "民营企业",
        "dwxzdm": "300",
        "dwhy": "信息传输、软件和信息技术服务业",
        "dwhydm": "I",
        "mobilePhone": "",
        "jtdh": "",
        "lxrdh": "",
        "lxrsj": "",
        "gzzwlb": "工程技术人员",
    }
    values.update(changes)
    for header, value in values.items():
        sheet.cell(2, header_columns[header]).value = value
    workbook.save(input_path)
    workbook.close()
    return input_path


def _reference_file(tmp_path: Path, record: ReferenceRecord) -> Path:
    path = tmp_path / "企业参考库.xlsx"
    write_reference_workbook(path, [record])
    return path


def _header_columns(sheet):
    return {cell.value: cell.column for cell in sheet[1]}


def _is_red(cell) -> bool:
    return cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb.endswith("FFE2E2")


def test_initial_and_final_real_samples_export_without_changing_input(tmp_path: Path):
    reference_path = PROJECT_ROOT / "企业参考库.xlsx"
    for stage, source_name in (("initial", "初审样例.xlsx"), ("final", "终审样例.xlsx")):
        input_path = tmp_path / source_name
        copy2(SAMPLES / source_name, input_path)
        before = _digest(input_path)

        outcome = audit_workbook(input_path, stage, reference_path)

        assert outcome.output_path.exists()
        assert _digest(input_path) == before
        workbook = load_workbook(outcome.output_path, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        workbook.close()
        assert headers[-len(RESULT_HEADERS) :] == RESULT_HEADERS
        assert outcome.review_stage.value == stage.upper()


def test_export_adds_results_and_highlights_actual_source_fields(tmp_path: Path):
    input_path = _prepare_input(
        tmp_path,
        "匿名导出.xlsx",
        dwmc="匿名错误企业",
        dwxzdm="WRONG-TYPE",
        dwhydm="WRONG-INDUSTRY",
        mobilePhone="13800138000",
        lxrdh="+86 138-0013-8000",
        gzzwlb="其他人员",
    )
    before = _digest(input_path)

    outcome = audit_workbook(input_path, "initial", _reference_file(tmp_path, _reference()))

    assert outcome.output_path.name == "匿名导出_初审审核结果.xlsx"
    assert _digest(input_path) == before
    workbook = load_workbook(outcome.output_path, data_only=True)
    sheet = workbook.active
    columns = _header_columns(sheet)
    assert sheet.cell(2, columns["审核阶段"]).value == "初审"
    assert sheet.cell(2, columns["总体审核结果"]).value == "有问题"
    assert sheet.cell(2, columns["企业信息校验"]).value.startswith("有问题：")
    assert sheet.cell(2, columns["电话校验"]).value.startswith("有问题：")
    assert sheet.cell(2, columns["单位行业校验"]).value.startswith("有问题：")
    assert sheet.cell(2, columns["单位性质校验"]).value.startswith("有问题：")
    assert sheet.cell(2, columns["工作职位校验"]).value.startswith("有问题：")
    assert "统一社会信用代码对应企业名称" in sheet.cell(2, columns["异常汇总"]).value
    assert _is_red(sheet["BB2"])
    assert _is_red(sheet["BC2"])
    assert _is_red(sheet["AR2"])
    assert _is_red(sheet["BM2"])
    assert _is_red(sheet["BD2"])
    assert _is_red(sheet["BE2"])
    assert _is_red(sheet["BF2"])
    assert _is_red(sheet["BG2"])
    assert _is_red(sheet["BR2"])
    assert not _is_red(sheet["BQ2"])
    workbook.close()


def test_export_marks_trusted_historical_match_as_normal(tmp_path: Path):
    input_path = _prepare_input(tmp_path, "匹配企业.xlsx")

    outcome = audit_workbook(input_path, "final", _reference_file(tmp_path, _reference()))

    workbook = load_workbook(outcome.output_path, data_only=True)
    sheet = workbook.active
    columns = _header_columns(sheet)
    assert sheet.cell(2, columns["审核阶段"]).value == "终审"
    assert sheet.cell(2, columns["总体审核结果"]).value == "正常"
    assert sheet.cell(2, columns["异常汇总"]).value in ("", None)
    workbook.close()


def test_export_surfaces_historical_reference_conflict(tmp_path: Path):
    input_path = _prepare_input(tmp_path, "冲突企业.xlsx")

    outcome = audit_workbook(
        input_path,
        "initial",
        _reference_file(tmp_path, replace(_reference(), reference_conflict=True, conflict_fields=("company_name",))),
    )

    workbook = load_workbook(outcome.output_path, data_only=True)
    sheet = workbook.active
    columns = _header_columns(sheet)
    assert "历史参考库中该统一社会信用代码存在冲突记录" in sheet.cell(2, columns["企业信息校验"]).value
    assert _is_red(sheet["BB2"])
    assert _is_red(sheet["BC2"])
    workbook.close()
