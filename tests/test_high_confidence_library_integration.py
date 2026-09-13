from pathlib import Path

from openpyxl import load_workbook
import pytest

from src.config import EXPECTED_REFERENCE_COUNT, FIELD_MAP
from src.reference import ReferenceLibrary, ReferenceLibraryIntegrityError
from src.rules import AuditEngine, ReviewStage


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OFFICIAL_REFERENCE = PROJECT_ROOT / "企业参考库.xlsx"
HIGH_CONFIDENCE_HEADERS = ["dwmc", "dwzzjgdm", "dwxz", "dwhy"]


def _row_for(record):
    return {
        FIELD_MAP["student_name"]: "测试学生",
        FIELD_MAP["student_id"]: "TEST-001",
        FIELD_MAP["company_name"]: record.company_name,
        FIELD_MAP["credit_code"]: record.credit_code,
        FIELD_MAP["company_type"]: record.company_type,
        FIELD_MAP["industry"]: record.industry,
        FIELD_MAP["job_category"]: "工程技术人员",
    }


def test_official_reference_is_the_validated_high_confidence_four_column_library():
    workbook = load_workbook(OFFICIAL_REFERENCE, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        assert [cell.value for cell in sheet[1]] == HIGH_CONFIDENCE_HEADERS
        assert sheet.max_row - 1 == EXPECTED_REFERENCE_COUNT
    finally:
        workbook.close()

    library = ReferenceLibrary.from_workbook(OFFICIAL_REFERENCE)
    assert len(library.by_credit_code) == EXPECTED_REFERENCE_COUNT
    assert all(record.reference_conflict is False for record in library.by_credit_code.values())
    assert all(record.credit_code and record.company_name and record.company_type and record.industry for record in library.by_credit_code.values())


def test_high_confidence_reference_drives_exact_company_audits():
    library = ReferenceLibrary.from_workbook(OFFICIAL_REFERENCE)
    reference = next(iter(library.by_credit_code.values()))
    engine = AuditEngine(library)

    matching = engine.audit(_row_for(reference), 2, ReviewStage.INITIAL)
    assert not any(issue.rule in {"credit_code_history_match", "company_name_match", "company_type_match", "industry_match"} for issue in matching.issues)

    wrong_name = _row_for(reference)
    wrong_name[FIELD_MAP["company_name"]] = reference.company_name + "测试"
    assert any(issue.rule == "company_name_match" for issue in engine.audit(wrong_name, 3, ReviewStage.INITIAL).issues)

    wrong_type = _row_for(reference)
    wrong_type[FIELD_MAP["company_type"]] = reference.company_type + "测试"
    assert any(issue.rule == "company_type_match" for issue in engine.audit(wrong_type, 4, ReviewStage.INITIAL).issues)

    wrong_industry = _row_for(reference)
    wrong_industry[FIELD_MAP["industry"]] = reference.industry + "测试"
    assert any(issue.rule == "industry_match" for issue in engine.audit(wrong_industry, 5, ReviewStage.INITIAL).issues)

    unknown = _row_for(reference)
    unknown[FIELD_MAP["credit_code"]] = "91UNKNOWN000000000"
    issues = engine.audit(unknown, 6, ReviewStage.INITIAL).issues
    assert any("未命中高可信企业参考库，请人工核实" in issue.message for issue in issues)


def test_five_high_confidence_companies_pass_the_company_reference_smoke_check():
    library = ReferenceLibrary.from_workbook(OFFICIAL_REFERENCE)
    engine = AuditEngine(library)
    company_rules = {"credit_code_history_match", "company_name_match", "company_type_match", "industry_match"}
    for row_number, reference in enumerate(list(library.by_credit_code.values())[:5], start=2):
        result = engine.audit(_row_for(reference), row_number, ReviewStage.INITIAL)
        assert not any(issue.rule in company_rules for issue in result.issues)


def test_strict_loader_rejects_the_retired_eight_column_reference_schema(tmp_path):
    workbook = load_workbook(OFFICIAL_REFERENCE, read_only=False)
    try:
        worksheet = workbook.active
        worksheet.delete_cols(1, worksheet.max_column)
        worksheet.append(["credit_code", "company_name", "company_type", "company_type_code", "industry", "industry_code", "reference_conflict", "conflict_fields"])
        worksheet.append(["91AAA", "甲有限公司", "民营企业", "", "制造业", "", False, ""])
        legacy_path = tmp_path / "legacy.xlsx"
        workbook.save(legacy_path)
    finally:
        workbook.close()
    with pytest.raises(ReferenceLibraryIntegrityError):
        ReferenceLibrary.from_workbook(legacy_path, require_high_confidence=True)
