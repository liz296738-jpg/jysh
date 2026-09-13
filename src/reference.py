"""Creation and lookup of a privacy-minimized, offline company reference library."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook

from .config import FIELD_MAP, TRUSTED_AUDIT_STATUS
from .normalization import normalize_company_name, normalize_credit_code, normalize_text


REFERENCE_HEADERS = [
    "credit_code",
    "company_name",
    "company_type",
    "company_type_code",
    "industry",
    "industry_code",
    "reference_conflict",
    "conflict_fields",
]

_REFERENCE_FIELD_KEYS = (
    "credit_code",
    "company_name",
    "company_type",
    "company_type_code",
    "industry",
    "industry_code",
)


def _clean_value(value: Any) -> str:
    return normalize_text(value)


@dataclass(frozen=True)
class ReferenceRecord:
    credit_code: str
    company_name: str
    company_type: str
    company_type_code: str
    industry: str
    industry_code: str
    reference_conflict: bool
    conflict_fields: tuple[str, ...]

    def as_row(self) -> list[object]:
        return [
            self.credit_code,
            self.company_name,
            self.company_type,
            self.company_type_code,
            self.industry,
            self.industry_code,
            self.reference_conflict,
            ", ".join(self.conflict_fields),
        ]


def _reference_values(row: dict[str, Any]) -> dict[str, str]:
    return {
        "credit_code": normalize_credit_code(row.get(FIELD_MAP["credit_code"])),
        "company_name": _clean_value(row.get(FIELD_MAP["company_name"])),
        "company_type": _clean_value(row.get(FIELD_MAP["company_type"])),
        "company_type_code": _clean_value(row.get(FIELD_MAP["company_type_code"])),
        "industry": _clean_value(row.get(FIELD_MAP["industry"])),
        "industry_code": _clean_value(row.get(FIELD_MAP["industry_code"])),
    }


def build_reference_records(rows: Iterable[dict[str, Any]]) -> list[ReferenceRecord]:
    """Build one record per safe company identity, retaining unresolved conflicts."""
    groups: dict[tuple[str, str], list[dict[str, str]]] = {}
    for row in rows:
        if _clean_value(row.get(FIELD_MAP["audit_status"])) != TRUSTED_AUDIT_STATUS:
            continue
        values = _reference_values(row)
        code = values["credit_code"]
        name = normalize_company_name(values["company_name"])
        if not code and not name:
            continue
        group_key = ("credit", code) if code else ("name", name)
        groups.setdefault(group_key, []).append(values)

    records: list[ReferenceRecord] = []
    for (_, _), members in groups.items():
        distinct = {
            key: {member[key] for member in members}
            for key in _REFERENCE_FIELD_KEYS
        }
        conflicts = tuple(key for key in _REFERENCE_FIELD_KEYS if len(distinct[key]) > 1)
        selected = {
            key: next(iter(distinct[key])) if len(distinct[key]) == 1 else ""
            for key in _REFERENCE_FIELD_KEYS
        }
        records.append(
            ReferenceRecord(
                **selected,
                reference_conflict=bool(conflicts),
                conflict_fields=conflicts,
            )
        )
    return sorted(records, key=lambda record: (record.credit_code == "", record.credit_code, record.company_name))


def write_reference_workbook(path: Path, records: Iterable[ReferenceRecord]) -> None:
    """Write only enterprise review fields; no student data enters this workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "企业参考库"
    sheet.append(REFERENCE_HEADERS)
    for record in records:
        sheet.append(record.as_row())
    workbook.save(path)
    workbook.close()


class ReferenceLibrary:
    """Load separate identity indexes and apply the credit-code-first lookup rule."""

    def __init__(self, records: Iterable[ReferenceRecord]) -> None:
        self.by_credit_code: dict[str, ReferenceRecord] = {}
        self.by_company_name: dict[str, ReferenceRecord] = {}
        for record in records:
            if record.credit_code:
                self.by_credit_code[normalize_credit_code(record.credit_code)] = record
            if not record.credit_code and record.company_name:
                self.by_company_name[normalize_company_name(record.company_name)] = record

    @classmethod
    def from_workbook(cls, path: Path) -> "ReferenceLibrary":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            missing = set(REFERENCE_HEADERS) - set(headers)
            if missing:
                raise ValueError(f"企业参考库缺少列：{', '.join(sorted(missing))}")
            records = []
            for values in sheet.iter_rows(min_row=2, values_only=True):
                row = dict(zip(headers, values))
                records.append(
                    ReferenceRecord(
                        credit_code=normalize_credit_code(row["credit_code"]),
                        company_name=_clean_value(row["company_name"]),
                        company_type=_clean_value(row["company_type"]),
                        company_type_code=_clean_value(row["company_type_code"]),
                        industry=_clean_value(row["industry"]),
                        industry_code=_clean_value(row["industry_code"]),
                        reference_conflict=bool(row["reference_conflict"]),
                        conflict_fields=tuple(
                            item.strip()
                            for item in str(row["conflict_fields"] or "").split(",")
                            if item.strip()
                        ),
                    )
                )
            return cls(records)
        finally:
            workbook.close()

    def find(self, credit_code: Any, company_name: Any) -> ReferenceRecord | None:
        normalized_code = normalize_credit_code(credit_code)
        if normalized_code:
            return self.by_credit_code.get(normalized_code)
        return self.by_company_name.get(normalize_company_name(company_name))
