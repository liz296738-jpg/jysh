"""Shared initial/final batch auditing and non-destructive Excel result export."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any
import warnings

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from .batch_rules import apply_batch_contact_phone_duplicates
from .config import FIELD_MAP
from .excel_parser import ExcelParser
from .normalization import normalize_phone
from .reference import ReferenceLibrary
from .rules import AuditIssue, AuditEngine, ReviewStage, StudentAuditResult


RESULT_HEADERS = [
    "审核阶段",
    "总体审核结果",
    "企业信息校验",
    "电话校验",
    "单位行业校验",
    "单位性质校验",
    "工作职位校验",
    "异常汇总",
]

_RED_FILL = PatternFill(fill_type="solid", fgColor="FFE2E2")
_STAGE_LABELS = {ReviewStage.INITIAL: "初审", ReviewStage.FINAL: "终审"}
_CATEGORY_RULES = {
    "企业信息校验": {
        "credit_code_history_match",
        "credit_code_reference_conflict",
        "company_name_match",
        "company_name_no_credit_history_match",
        "company_name_reference_conflict",
    },
    "电话校验": {
        "company_contact_not_student_phone",
        "company_contact_not_family_phone",
        "batch_company_contact_phone_duplicate",
    },
    "单位行业校验": {"industry_match"},
    "单位性质校验": {"company_type_match"},
    "工作职位校验": {"job_category_not_other"},
}


@dataclass(frozen=True)
class AuditWorkbookOutcome:
    output_path: Path
    review_stage: ReviewStage
    results: tuple[StudentAuditResult, ...]


def _parse_stage(review_stage: str | ReviewStage) -> ReviewStage:
    if isinstance(review_stage, ReviewStage):
        return review_stage
    try:
        return ReviewStage(review_stage.upper())
    except (AttributeError, ValueError) as error:
        raise ValueError("review_stage 必须为 'initial' 或 'final'") from error


def _result_text(issues: list[AuditIssue], rules: set[str]) -> str:
    messages = [issue.message for issue in issues if issue.rule in rules]
    return "正常" if not messages else f"有问题：{'；'.join(messages)}"


def _matching_contact_headers(row: dict[str, Any], target_phone: str) -> set[str]:
    normalized_target = normalize_phone(target_phone)
    if not normalized_target:
        return set()
    return {
        FIELD_MAP[field_key]
        for field_key in ("company_phone", "company_mobile")
        if normalize_phone(row.get(FIELD_MAP[field_key])) == normalized_target
    }


def _highlight_headers(issue: AuditIssue, row: dict[str, Any]) -> set[str]:
    if issue.rule in _CATEGORY_RULES["企业信息校验"]:
        return {FIELD_MAP["company_name"], FIELD_MAP["credit_code"]}
    if issue.rule == "industry_match":
        return {FIELD_MAP["industry"], FIELD_MAP["industry_code"]}
    if issue.rule == "company_type_match":
        return {FIELD_MAP["company_type"], FIELD_MAP["company_type_code"]}
    if issue.rule == "job_category_not_other":
        return {FIELD_MAP["job_category"]}
    if issue.rule == "company_contact_not_student_phone":
        return {FIELD_MAP["student_phone"]} | _matching_contact_headers(row, issue.student_value)
    if issue.rule == "company_contact_not_family_phone":
        return {FIELD_MAP["family_phone"]} | _matching_contact_headers(row, issue.student_value)
    if issue.rule == "batch_company_contact_phone_duplicate":
        return {FIELD_MAP["company_phone"], FIELD_MAP["company_mobile"]}
    return set(issue.highlight_fields)


def _append_result_columns(sheet, result: StudentAuditResult, row: dict[str, Any], row_number: int, columns: dict[str, int], stage: ReviewStage) -> None:
    values = [
        _STAGE_LABELS[stage],
        "正常" if result.is_ok else "有问题",
        _result_text(result.issues, _CATEGORY_RULES["企业信息校验"]),
        _result_text(result.issues, _CATEGORY_RULES["电话校验"]),
        _result_text(result.issues, _CATEGORY_RULES["单位行业校验"]),
        _result_text(result.issues, _CATEGORY_RULES["单位性质校验"]),
        _result_text(result.issues, _CATEGORY_RULES["工作职位校验"]),
        "；".join(issue.message for issue in result.issues),
    ]
    first_result_column = max(columns.values()) + 1
    for offset, value in enumerate(values):
        sheet.cell(row_number, first_result_column + offset).value = value
    for issue in result.issues:
        for header in _highlight_headers(issue, row):
            if header in columns:
                sheet.cell(row_number, columns[header]).fill = _RED_FILL


def audit_workbook(
    input_path: str | Path,
    review_stage: str | ReviewStage,
    reference_path: str | Path | None = None,
    require_high_confidence_reference: bool = False,
) -> AuditWorkbookOutcome:
    """Audit a single export without changing it and write a sibling result workbook."""
    source_path = Path(input_path)
    stage = _parse_stage(review_stage)
    project_root = Path(__file__).resolve().parents[1]
    library_path = Path(reference_path) if reference_path else project_root / "企业参考库.xlsx"
    parsed = ExcelParser(FIELD_MAP).read(source_path)
    engine = AuditEngine(
        ReferenceLibrary.from_workbook(
            library_path, require_high_confidence=require_high_confidence_reference
        )
    )
    results = [
        engine.audit(row, index + 2, stage)
        for index, row in enumerate(parsed.rows)
    ]
    apply_batch_contact_phone_duplicates(results, parsed.rows)

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style.*",
            category=UserWarning,
            module="openpyxl.styles.stylesheet",
        )
        workbook = load_workbook(source_path)
    try:
        sheet = workbook.active
        columns = {cell.value: cell.column for cell in sheet[1] if cell.value}
        first_result_column = sheet.max_column + 1
        for offset, header in enumerate(RESULT_HEADERS):
            sheet.cell(1, first_result_column + offset).value = header
        for result, row in zip(results, parsed.rows):
            _append_result_columns(sheet, result, row, result.row_number, columns, stage)
        output_path = source_path.with_name(
            f"{source_path.stem}_{_STAGE_LABELS[stage]}审核结果{source_path.suffix}"
        )
        workbook.save(output_path)
    finally:
        workbook.close()
    return AuditWorkbookOutcome(output_path, stage, tuple(results))
