"""Header-driven parsing for the school's employment Excel exports."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping
import warnings

from openpyxl import load_workbook


class MissingRequiredColumnsError(ValueError):
    """Raised when a selected export cannot support the configured review fields."""


@dataclass(frozen=True)
class ParsedWorkbook:
    sheet_name: str
    headers: frozenset[str]
    row_count: int
    rows: tuple[dict[str, Any], ...]


class ExcelParser:
    """Read a workbook by header names, never by fixed column positions."""

    def __init__(self, field_map: Mapping[str, str]) -> None:
        self._required_headers = frozenset(field_map.values())

    def read(self, path: Path) -> ParsedWorkbook:
        with warnings.catch_warnings():
            warnings.filterwarnings(
                "ignore",
                message="Workbook contains no default style.*",
                category=UserWarning,
                module="openpyxl.styles.stylesheet",
            )
            workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = workbook.active
            header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers = tuple(value.strip() if isinstance(value, str) else value for value in header_values)
            header_set = frozenset(value for value in headers if isinstance(value, str) and value)
            missing = sorted(self._required_headers - header_set)
            if missing:
                raise MissingRequiredColumnsError(
                    f"Excel 缺少关键列：{', '.join(missing)}"
                )

            rows = tuple(
                {
                    header: value
                    for header, value in zip(headers, values)
                    if isinstance(header, str) and header
                }
                for values in sheet.iter_rows(min_row=2, values_only=True)
            )
            return ParsedWorkbook(
                sheet_name=sheet.title,
                headers=header_set,
                row_count=len(rows),
                rows=rows,
            )
        finally:
            workbook.close()
