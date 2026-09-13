"""Safe, local management of the static Excel enterprise reference library."""

from __future__ import annotations

import json
import os
import re
import shutil
import threading
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .config import HIGH_CONFIDENCE_REFERENCE_HEADERS
from .normalization import normalize_company_name, normalize_credit_code
from .reference import ReferenceLibrary, ReferenceLibraryIntegrityError


_WRITE_LOCK = threading.RLock()
_CODE_PATTERN = re.compile(r"^[0-9A-Z]{18}$")


class ReferenceManagementError(ValueError):
    pass


def _display(value: Any) -> str:
    return str(value or "").strip()


def _missing(value: str) -> bool:
    return not value or value == "30"


class ReferenceManager:
    def __init__(self, path: Path, backup_root: Path, log_path: Path, backup_limit: int = 30) -> None:
        self.path, self.backup_root, self.log_path, self.backup_limit = Path(path), Path(backup_root), Path(log_path), backup_limit

    def records(self) -> list[dict[str, str]]:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            headers = tuple(cell.value for cell in sheet[1])
            if headers != HIGH_CONFIDENCE_REFERENCE_HEADERS:
                raise ReferenceLibraryIntegrityError("企业高可信参考库完整性校验失败")
            rows = [dict(zip(headers, (_display(value) for value in values))) for values in sheet.iter_rows(min_row=2, values_only=True)]
        finally:
            workbook.close()
        ReferenceLibrary.from_workbook(self.path, require_high_confidence=True)
        return rows

    def choices(self) -> tuple[list[str], list[str]]:
        rows = self.records()
        return sorted({row["dwxz"] for row in rows}), sorted({row["dwhy"] for row in rows})

    def find(self, code: str) -> dict[str, str] | None:
        normalized = normalize_credit_code(code)
        return next((row for row in self.records() if normalize_credit_code(row["dwzzjgdm"]) == normalized), None)

    def save(self, action: str, submitted: dict[str, Any], original_code: str | None = None) -> dict[str, str]:
        with _WRITE_LOCK:
            rows = self.records()
            record = self._validate_submission(submitted, rows, original_code)
            before = None
            if action == "EDIT":
                old = normalize_credit_code(original_code)
                for index, row in enumerate(rows):
                    if normalize_credit_code(row["dwzzjgdm"]) == old:
                        before = dict(row)
                        rows[index] = record
                        break
                else:
                    raise ReferenceManagementError("未找到要编辑的企业记录。")
            elif action == "ADD":
                rows.append(record)
            else:
                raise ReferenceManagementError("不支持的参考库操作。")
            self._write_atomically(rows)
            self._append_log(action, record["dwzzjgdm"], before, record)
            return record

    def _validate_submission(self, submitted: dict[str, Any], rows: list[dict[str, str]], original_code: str | None) -> dict[str, str]:
        record = {field: _display(submitted.get(field)) for field in HIGH_CONFIDENCE_REFERENCE_HEADERS}
        record["dwzzjgdm"] = normalize_credit_code(record["dwzzjgdm"])
        labels = {"dwmc": "单位名称", "dwzzjgdm": "统一社会信用代码", "dwxz": "单位性质", "dwhy": "单位行业"}
        for field, label in labels.items():
            if _missing(record[field]):
                if record[field] == "30":
                    raise ReferenceManagementError("30 为历史缺失占位值，不能写入高可信参考库。")
                raise ReferenceManagementError(f"{label}不能为空。")
        if not _CODE_PATTERN.fullmatch(record["dwzzjgdm"]):
            raise ReferenceManagementError("统一社会信用代码必须为 18 位数字或大写英文字母。")
        original = normalize_credit_code(original_code) if original_code else None
        for row in rows:
            code = normalize_credit_code(row["dwzzjgdm"])
            if code != original and code == record["dwzzjgdm"]:
                raise ReferenceManagementError("统一社会信用代码已存在。")
            if code != original and normalize_company_name(row["dwmc"]) == normalize_company_name(record["dwmc"]):
                raise ReferenceManagementError("单位名称已对应其他统一社会信用代码，请人工核实。")
        return record

    def _write_atomically(self, rows: list[dict[str, str]]) -> None:
        temporary = self.path.with_suffix(".tmp.xlsx")
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.append(HIGH_CONFIDENCE_REFERENCE_HEADERS)
            for row in rows:
                sheet.append([row[field] for field in HIGH_CONFIDENCE_REFERENCE_HEADERS])
            workbook.save(temporary)
        finally:
            workbook.close()
        try:
            ReferenceLibrary.from_workbook(temporary, require_high_confidence=True)
            self.backup_root.mkdir(parents=True, exist_ok=True)
            backup = self.backup_root / f"企业参考库_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
            shutil.copy2(self.path, backup)
            os.replace(temporary, self.path)
            ReferenceLibrary.from_workbook(self.path, require_high_confidence=True)
            backups = sorted(self.backup_root.glob("*.xlsx"), key=lambda item: item.stat().st_mtime, reverse=True)
            for stale in backups[self.backup_limit:]:
                stale.unlink()
        except Exception:
            if temporary.exists():
                temporary.unlink()
            raise

    def _append_log(self, action: str, code: str, before: dict[str, str] | None, after: dict[str, str]) -> None:
        self.log_path.parent.mkdir(parents=True, exist_ok=True)
        entry = {"timestamp": datetime.now().isoformat(timespec="seconds"), "action": action, "credit_code": code, "before": before, "after": after}
        with self.log_path.open("a", encoding="utf-8") as handle:
            handle.write(json.dumps(entry, ensure_ascii=False) + "\n")
